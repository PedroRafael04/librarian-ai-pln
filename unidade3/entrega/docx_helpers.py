"""Primitivas de formatacao do documento de entrega (python-docx).

O estilo segue o modelo de entrega da disciplina: A4, corpo em Times New Roman,
titulos de secao em azul-escuro, tabelas em Calibri com cabecalho sombreado,
legendas de figura em Calibri italico e trechos de codigo em Consolas sobre
fundo cinza.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (reexportado para o gerador)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Times New Roman"
TABLE_FONT = "Calibri"
CODE_FONT = "Consolas"
HEADING_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
HEADER_FILL = "DCE6F0"
CODE_FILL = "F3F3F3"
BODY_SIZE = 10.5
TABLE_SIZE = 8.5


def _set_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)


def _shade(element, fill: str) -> None:
    """Sombreamento de celula (``tcPr``) ou de paragrafo (``pPr``)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


# Marcacao inline minima: **negrito**, *italico*, `codigo`.
_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


def add_runs(paragraph, text: str, font: str = BODY_FONT, size: float = BODY_SIZE,
             bold: bool = False, italic: bool = False) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        b, i, f = bold, italic, font
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, f = part[1:-1], CODE_FONT
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        run = paragraph.add_run(part)
        _set_font(run, f, size - 1 if f == CODE_FONT else size)
        run.bold, run.italic = b, i


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        normal = self.doc.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(BODY_SIZE)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(4), 1.08
        self.figures = 0
        self.tables = 0
        self._page_numbers()

    @property
    def text_width(self):
        sec = self.doc.sections[0]
        return sec.page_width - sec.left_margin - sec.right_margin

    def _page_numbers(self) -> None:
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        _set_font(run, BODY_FONT, 9)
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._element.append(el)

    # -- blocos de texto ---------------------------------------------------

    def centered(self, text: str, bold: bool = False, italic: bool = False,
                 size: float = BODY_SIZE, after: float = 2) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        add_runs(p, text, size=size, bold=bold, italic=italic)

    def heading(self, text: str, size: float = 13, before: float = 14) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        _set_font(run, BODY_FONT, size)
        run.bold = True
        run.font.color.rgb = HEADING_COLOR

    def subheading(self, text: str) -> None:
        self.heading(text, size=11.5, before=8)

    def para(self, text: str, after: float = 4, keep: bool = False):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_with_next = keep
        add_runs(p, text)
        return p

    def bullets(self, items: list[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, item)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- figuras, tabelas e codigo -----------------------------------------

    def figure(self, path: Path, caption: str, width_cm: float = 16.6) -> None:
        self.figures += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(8)
        add_runs(c, f"Figura {self.figures} — {caption}", font=TABLE_FONT, size=8.5, italic=True)

    def table(self, caption: str, headers: list[str], rows: list[list], widths: list[float] | None = None,
              align_right: set[int] | None = None, size: float = TABLE_SIZE,
              numbered: bool = True) -> None:
        if numbered:
            self.tables += 1
            caption = f"Tabela {self.tables} — {caption}"
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(3)
        add_runs(cap, caption, font=TABLE_FONT, size=9.5, bold=True)
        # Tabelas curtas ficam inteiras numa pagina (keep-with-next em cada linha).
        keep_rows = len(rows) <= 20

        align_right = align_right or set()
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = widths is None
        for r, values in enumerate([headers, *rows]):
            row = t.rows[r]
            # Linhas nao quebram entre paginas; o cabecalho se repete.
            trpr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trpr.append(cant)
            if r == 0:
                hdr = OxmlElement("w:tblHeader")
                trpr.append(hdr)
            for c, value in enumerate(values):
                cell = row.cells[c]
                if widths:
                    cell.width = Cm(widths[c])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = keep_rows and r < len(rows)
                if c in align_right and r > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                add_runs(p, str(value), font=TABLE_FONT, size=size, bold=(r == 0))
                if r == 0:
                    _shade(cell._tc.get_or_add_tcPr(), HEADER_FILL)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def code(self, title: str, source: str) -> None:
        cap = self.doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(2)
        add_runs(cap, title, font=TABLE_FONT, size=9.5, bold=True)
        lines = source.strip("\n").split("\n")
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            pf.keep_together = True
            if i < len(lines) - 1:
                pf.keep_with_next = len(lines) <= 30
            _shade(p._p.get_or_add_pPr(), CODE_FILL)
            run = p.add_run(line if line else " ")
            _set_font(run, CODE_FONT, 7.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
